"""Strict static-document subset, not antivirus or a general document renderer.

This module is only launched by validation.py. POSIX resource limits are required;
unsupported platforms fail closed. No source URLs, macros or formulas are executed.
"""

import csv
import io
import json
import re
import sys
import zipfile

from investment_os.ingestion.validation import MAX_BYTES, POLICY_VERSION


class Rejected(ValueError):
    pass


def require(condition, reason):
    if not condition:
        raise Rejected(reason)


def formula(value: str) -> bool:
    value = value.lstrip()
    return value.startswith(("=", "+", "@")) or (
        value.startswith("-") and not re.fullmatch(r"-\d+(?:\.\d+)?(?:[eE][+-]?\d+)?", value)
    )


def inspect_csv(data):
    require(not data.startswith((b"%PDF-", b"PK", b"\xd0\xcf")), "magic_mismatch")
    text = data.decode("utf-8-sig", errors="strict")
    require(not any(ord(c) < 32 and c not in "\r\n\t" for c in text), "binary_csv")
    csv.field_size_limit(10_000)
    rows = csv.reader(io.StringIO(text, newline=""), strict=True)
    count, width = 0, None
    for row in rows:
        count += 1
        require(count <= 10_000 and 1 <= len(row) <= 100, "csv_dimensions")
        if width is None:
            width = len(row)
        require(len(row) == width, "ragged_csv")
        require(not any(formula(cell) for cell in row), "formula_cell")
    require(count > 0, "empty_csv")
    return {"rows": count, "columns": width}


def inspect_xlsx(data):
    from defusedxml.ElementTree import fromstring
    from openpyxl import load_workbook

    require(data.startswith(b"PK\x03\x04"), "magic_mismatch")
    # A conservative OOXML allowlist: static values and styles, no drawings/links/macros.
    allowed = re.compile(
        r"(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|"
        r"xl/(?:workbook|styles|sharedStrings)\.xml|xl/_rels/workbook\.xml\.rels|"
        r"xl/theme/theme\d+\.xml|xl/worksheets/sheet\d+\.xml)"
    )
    required = {
        "[Content_Types].xml",
        "_rels/.rels",
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
    }
    denied_tags = {
        "formula",
        "formula1",
        "formula2",
        "definedName",
        "conditionalFormatting",
        "dataValidation",
        "dataValidations",
        "f",
        "externalLink",
        "externalReference",
        "externalReferences",
        "hyperlink",
        "hyperlinks",
        "oleObject",
        "oleObjects",
        "ddeLink",
        "extLst",
    }
    total = 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        names = [entry.filename for entry in entries]
        require(0 < len(entries) <= 100 and len(names) == len(set(names)), "archive_entries")
        require(required.issubset(names), "xlsx_package_incomplete")
        for entry in entries:
            require(bool(allowed.fullmatch(entry.filename)), "unsupported_archive_part")
            require(not entry.flag_bits & 1, "encrypted_archive")
            require(
                entry.compress_type in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED},
                "archive_compression",
            )
            require((entry.external_attr >> 16) & 0o170000 != 0o120000, "archive_symlink")
            total += entry.file_size
            require(entry.file_size <= 1_000_000 and total <= 4_000_000, "archive_expanded_size")
            require(entry.file_size <= max(entry.compress_size, 1) * 100, "archive_expansion_ratio")
            # read() verifies CRC; never extract to disk. XML entities/DTD are disabled.
            tree = fromstring(
                archive.read(entry), forbid_dtd=True, forbid_entities=True, forbid_external=True
            )
            for node in tree.iter():
                local = node.tag.split("}")[-1]
                require(local not in denied_tags, "active_spreadsheet_content")
                if local == "Relationship":
                    require(
                        node.get("TargetMode", "Internal") == "Internal", "external_relationship"
                    )
                    target = node.get("Target", "")
                    require(
                        ":" not in target and "\\" not in target and ".." not in target.split("/"),
                        "relationship_path",
                    )
                if local in {"Override", "Default"}:
                    media = node.get("ContentType", "").lower()
                    require(
                        not any(word in media for word in ("macro", "vba", "ole", "activex")),
                        "active_content_type",
                    )
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        require(1 <= len(workbook.worksheets) <= 10, "sheet_count")
        cells = 0
        for sheet in workbook.worksheets:
            require(
                (sheet.max_row or 0) <= 10_000 and (sheet.max_column or 0) <= 100,
                "sheet_dimensions",
            )
            for row in sheet.iter_rows():
                cells += len(row)
                require(cells <= 50_000, "cell_count")
                for cell in row:
                    require(cell.data_type != "f", "formula_cell")
                    if isinstance(cell.value, str):
                        require(
                            len(cell.value) <= 10_000 and not formula(cell.value),
                            "unsafe_string_cell",
                        )
        return {"sheets": len(workbook.worksheets), "cells": cells}
    finally:
        workbook.close()


def inspect_pdf(data):
    from pypdf import PdfReader
    from pypdf.generic import ArrayObject, DictionaryObject, IndirectObject

    require(bool(re.match(rb"%PDF-1\.[0-7][\r\n]", data)), "magic_mismatch")
    require(data.rstrip().endswith(b"%%EOF"), "pdf_trailing_or_incomplete")
    reader = PdfReader(io.BytesIO(data), strict=True)
    require(not reader.is_encrypted, "encrypted_pdf")
    forbidden = {
        "/Annots",
        "/F",
        "/FFilter",
        "/FDecodeParms",
        "/Movie",
        "/Sound",
        "/3D",
        "/Screen",
        "/AA",
        "/OpenAction",
        "/JS",
        "/JavaScript",
        "/URI",
        "/Launch",
        "/EmbeddedFiles",
        "/EmbeddedFile",
        "/Filespec",
        "/AcroForm",
        "/XFA",
        "/RichMedia",
        "/Rendition",
        "/SubmitForm",
        "/GoToR",
        "/GoToE",
    }
    seen, visited = set(), 0

    def visit(value, depth=0):
        nonlocal visited
        visited += 1
        require(visited <= 50_000 and depth <= 50, "pdf_graph_limit")
        if isinstance(value, IndirectObject):
            key = (value.idnum, value.generation)
            if key in seen:
                return
            seen.add(key)
            require(len(seen) <= 5000, "pdf_object_limit")
            visit(value.get_object(), depth + 1)
        elif isinstance(value, DictionaryObject):
            require(not any(str(key) in forbidden for key in value), "active_pdf_content")
            for item in value.values():
                visit(item, depth + 1)
        elif isinstance(value, ArrayObject):
            for item in value:
                visit(item, depth + 1)
        elif isinstance(value, str):
            require(value not in forbidden, "active_pdf_content")

    visit(reader.trailer)
    # Inspect unreferenced xref objects as well; compressed object expansion is resource-limited.
    objects = {
        (number, generation)
        for generation, entries in reader.xref.items()
        if generation != 65535
        for number in entries
        if number
    }
    objects.update((number, 0) for number in reader.xref_objStm)
    require(len(objects) <= 5000, "pdf_object_limit")
    for number, generation in sorted(objects):
        visit(IndirectObject(number, generation, reader))
    require(1 <= len(reader.pages) <= 100, "pdf_page_count")
    return {"pages": len(reader.pages), "objects": len(objects)}


def main():
    try:
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (512 * 1024 * 1024,) * 2)
        resource.setrlimit(resource.RLIMIT_CPU, (4, 4))
        resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))
        resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
        data = sys.stdin.buffer.read(MAX_BYTES + 1)
        require(0 < len(data) <= MAX_BYTES, "size_limit")
        details = {"csv": inspect_csv, "xlsx": inspect_xlsx, "pdf": inspect_pdf}[sys.argv[1]](data)
        result = {"accepted": True, "details": details}
    except Rejected as exc:
        result = {"accepted": False, "reason": str(exc)}
    except Exception:
        # Never return parser messages containing document text, URLs or local paths.
        result = {"accepted": False, "reason": "malformed_or_unsupported_document"}
    print(json.dumps(dict(result, policy=POLICY_VERSION)))


if __name__ == "__main__":
    main()
